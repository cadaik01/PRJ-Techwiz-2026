from datetime import timedelta
from io import BytesIO

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from orders.models import OrderStatus
from system.excel import XLSX_CONTENT_TYPE
from system.models import AuditAction, AuditLog
from system.reports import MAX_RANGE_DAYS

SUMMARY_URL = "admin-report-summary"
EXPORT_URL = "admin-report-export"


@pytest.fixture
def window():
    today = timezone.localdate()
    return str(today), str(today + timedelta(days=3))


@pytest.mark.django_db
def test_revenue_counts_only_completed_orders(admin_client, market, make_order, window):
    today = timezone.localdate()
    make_order(pickup_date=today, status=OrderStatus.COMPLETED, total="10.00")
    make_order(pickup_date=today + timedelta(days=1), status=OrderStatus.COMPLETED, total="5.50")
    make_order(pickup_date=today, status=OrderStatus.PLACED, total="99.00")
    make_order(pickup_date=today, status=OrderStatus.CANCELLED, total="99.00")

    date_from, date_to = window
    data = admin_client.get(reverse(SUMMARY_URL), {"from": date_from, "to": date_to}).data["data"]

    assert len(data["revenue_by_market"]) == 1
    row = data["revenue_by_market"][0]
    assert row["market_name"] == "Central Market"
    assert row["completed_orders"] == 2
    # Money is a decimal string, never a float (D-020).
    assert row["revenue"] == "15.50"


@pytest.mark.django_db
def test_orders_by_status_covers_every_status_in_the_range(admin_client, market, make_order, window):
    today = timezone.localdate()
    make_order(pickup_date=today, status=OrderStatus.COMPLETED)
    make_order(pickup_date=today, status=OrderStatus.NO_SHOW)

    date_from, date_to = window
    rows = admin_client.get(reverse(SUMMARY_URL), {"from": date_from, "to": date_to}).data["data"][
        "orders_by_status"
    ]

    counts = {row["status"]: row["count"] for row in rows}
    assert [row["status"] for row in rows] == list(OrderStatus.values)
    assert counts["COMPLETED"] == 1
    assert counts["NO_SHOW"] == 1
    assert counts["EXPIRED"] == 0


@pytest.mark.django_db
def test_the_range_filters_on_the_pickup_date(admin_client, market, make_order):
    today = timezone.localdate()
    make_order(pickup_date=today, status=OrderStatus.COMPLETED, total="10.00")
    make_order(pickup_date=today + timedelta(days=30), status=OrderStatus.COMPLETED, total="77.00")

    data = admin_client.get(
        reverse(SUMMARY_URL), {"from": str(today), "to": str(today + timedelta(days=1))}
    ).data["data"]

    assert data["revenue_by_market"][0]["revenue"] == "10.00"


@pytest.mark.django_db
def test_both_range_bounds_are_inclusive(admin_client, market, make_order):
    today = timezone.localdate()
    make_order(pickup_date=today, status=OrderStatus.COMPLETED, total="1.00")
    make_order(pickup_date=today + timedelta(days=2), status=OrderStatus.COMPLETED, total="2.00")

    data = admin_client.get(
        reverse(SUMMARY_URL), {"from": str(today), "to": str(today + timedelta(days=2))}
    ).data["data"]

    assert data["revenue_by_market"][0]["completed_orders"] == 2
    assert data["revenue_by_market"][0]["revenue"] == "3.00"


@pytest.mark.django_db
def test_market_id_narrows_the_report(admin_client, market, other_market, make_order, window):
    today = timezone.localdate()
    make_order(pickup_date=today, status=OrderStatus.COMPLETED, total="10.00")
    make_order(
        pickup_date=today, status=OrderStatus.COMPLETED, total="20.00", order_market=other_market
    )

    date_from, date_to = window
    data = admin_client.get(
        reverse(SUMMARY_URL), {"from": date_from, "to": date_to, "market_id": other_market.id}
    ).data["data"]

    assert len(data["revenue_by_market"]) == 1
    assert data["revenue_by_market"][0]["market_id"] == other_market.id
    assert data["revenue_by_market"][0]["revenue"] == "20.00"


@pytest.mark.django_db
def test_top_farmers_reports_revenue_and_rating(
    admin_client, market, approved_farmer, make_order, window
):
    from reviews.models import FarmerReview

    today = timezone.localdate()
    first = make_order(pickup_date=today, status=OrderStatus.COMPLETED, total="10.00")
    make_order(pickup_date=today, status=OrderStatus.COMPLETED, total="4.00")
    FarmerReview.objects.create(order=first, rating=4)

    date_from, date_to = window
    rows = admin_client.get(reverse(SUMMARY_URL), {"from": date_from, "to": date_to}).data["data"][
        "top_farmers"
    ]

    assert len(rows) == 1
    row = rows[0]
    assert row["farmer_id"] == approved_farmer.user_id
    assert row["stall_name"] == "Test Stall"
    # The rating join must not multiply the order rows.
    assert row["completed_orders"] == 2
    assert row["revenue"] == "14.00"
    assert row["rating_avg"] == 4.0


@pytest.mark.django_db
def test_a_farmer_without_reviews_has_a_null_rating(admin_client, market, make_order, window):
    make_order(pickup_date=timezone.localdate(), status=OrderStatus.COMPLETED, total="10.00")

    date_from, date_to = window
    rows = admin_client.get(reverse(SUMMARY_URL), {"from": date_from, "to": date_to}).data["data"][
        "top_farmers"
    ]

    assert rows[0]["rating_avg"] is None


@pytest.mark.django_db
def test_hidden_reviews_are_left_out_of_the_rating(
    admin_client, market, approved_farmer, make_order, window
):
    from reviews.models import FarmerReview

    order = make_order(pickup_date=timezone.localdate(), status=OrderStatus.COMPLETED)
    FarmerReview.objects.create(order=order, rating=1, is_hidden_by_admin=True)

    date_from, date_to = window
    rows = admin_client.get(reverse(SUMMARY_URL), {"from": date_from, "to": date_to}).data["data"][
        "top_farmers"
    ]

    assert rows[0]["rating_avg"] is None


@pytest.mark.django_db
def test_an_empty_range_returns_empty_blocks(admin_client, market, window):
    date_from, date_to = window
    data = admin_client.get(reverse(SUMMARY_URL), {"from": date_from, "to": date_to}).data["data"]

    assert data["revenue_by_market"] == []
    assert data["top_farmers"] == []
    assert all(row["count"] == 0 for row in data["orders_by_status"])


@pytest.mark.django_db
def test_both_dates_are_required(admin_client):
    today = str(timezone.localdate())

    for params, missing in (({}, "from"), ({"from": today}, "to"), ({"to": today}, "from")):
        response = admin_client.get(reverse(SUMMARY_URL), params)
        assert response.status_code == 400
        assert response.data["code"] == "VALIDATION_ERROR"
        assert missing in response.data["errors"]


@pytest.mark.django_db
def test_a_malformed_date_is_a_field_error(admin_client):
    response = admin_client.get(
        reverse(SUMMARY_URL), {"from": "25-09-2026", "to": str(timezone.localdate())}
    )

    assert response.status_code == 400
    assert "from" in response.data["errors"]


@pytest.mark.django_db
def test_end_before_start_is_rejected(admin_client):
    today = timezone.localdate()

    response = admin_client.get(
        reverse(SUMMARY_URL), {"from": str(today), "to": str(today - timedelta(days=1))}
    )

    assert response.status_code == 400
    assert "to" in response.data["errors"]


@pytest.mark.django_db
def test_a_range_longer_than_the_cap_is_rejected(admin_client):
    today = timezone.localdate()

    ok = admin_client.get(
        reverse(SUMMARY_URL),
        {"from": str(today), "to": str(today + timedelta(days=MAX_RANGE_DAYS - 1))},
    )
    assert ok.status_code == 200

    too_long = admin_client.get(
        reverse(SUMMARY_URL),
        {"from": str(today), "to": str(today + timedelta(days=MAX_RANGE_DAYS))},
    )
    assert too_long.status_code == 400
    assert "to" in too_long.data["errors"]


@pytest.mark.django_db
def test_export_returns_a_three_sheet_workbook_and_audits(
    admin_client, market, make_order, window, admin_user
):
    today = timezone.localdate()
    make_order(pickup_date=today, status=OrderStatus.COMPLETED, total="12.50")

    date_from, date_to = window
    response = admin_client.get(reverse(EXPORT_URL), {"from": date_from, "to": date_to})

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    assert "attachment" in response["Content-Disposition"]
    assert ".xlsx" in response["Content-Disposition"]

    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Orders by status", "Revenue by market", "Top farmers"]

    revenue = workbook["Revenue by market"]
    assert [cell.value for cell in revenue[1]] == ["Market", "Completed orders", "Revenue (USD)"]
    assert revenue.cell(row=2, column=1).value == "Central Market"
    # Revenue is written as a number so Excel can sum the column.
    assert revenue.cell(row=2, column=3).value == 12.5

    entry = AuditLog.objects.get(action=AuditAction.EXPORT_DATA)
    assert entry.user == admin_user
    assert entry.details["from"] == date_from
    assert entry.details["to"] == date_to


@pytest.mark.django_db
def test_export_errors_still_use_the_json_envelope(admin_client):
    response = admin_client.get(reverse(EXPORT_URL), {"from": "not-a-date"})

    assert response.status_code == 400
    assert response.data["code"] == "VALIDATION_ERROR"
    assert not AuditLog.objects.filter(action=AuditAction.EXPORT_DATA).exists()


@pytest.mark.django_db
def test_customer_cannot_reach_the_reports(customer_client, window):
    date_from, date_to = window
    params = {"from": date_from, "to": date_to}

    assert customer_client.get(reverse(SUMMARY_URL), params).status_code == 403
    assert customer_client.get(reverse(EXPORT_URL), params).status_code == 403
