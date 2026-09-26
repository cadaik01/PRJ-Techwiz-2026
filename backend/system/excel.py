from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SHEETS = (
    ("Orders by status", ["Status", "Orders"], ("status", "count")),
    (
        "Revenue by market",
        ["Market", "Completed orders", "Revenue (USD)"],
        ("market_name", "completed_orders", "revenue"),
    ),
    (
        "Top farmers",
        ["Stall", "Completed orders", "Revenue (USD)", "Average rating"],
        ("stall_name", "completed_orders", "revenue", "rating_avg"),
    ),
)
SHEET_KEYS = ("orders_by_status", "revenue_by_market", "top_farmers")
MIN_WIDTH = 12
PADDING = 2


def report_filename(*, date_from: date, date_to: date) -> str:
    return f"marketlink-report-{date_from:%Y%m%d}-{date_to:%Y%m%d}.xlsx"


def _autosize(sheet, columns: int) -> None:
    for index in range(1, columns + 1):
        longest = max(
            (len(str(cell.value)) for cell in sheet[get_column_letter(index)] if cell.value),
            default=0,
        )
        sheet.column_dimensions[get_column_letter(index)].width = max(MIN_WIDTH, longest + PADDING)


def _keep_text_as_text(cells) -> None:
    # openpyxl stores any string starting with "=" as a formula. Names come from users (a
    # farmer picks the stall name), so every text cell is forced back to plain text: a name
    # like "=HYPERLINK(...)" is shown, never evaluated, when an admin opens the file.
    for cell in cells:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def build_report_workbook(summary: dict) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for key, (title, headers, fields) in zip(SHEET_KEYS, SHEETS, strict=True):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in summary[key]:
            # Money is a decimal string in JSON (D-020); Excel needs a number to sum it.
            sheet.append(
                [float(row[f]) if f == "revenue" else row[f] for f in fields]
            )
            _keep_text_as_text(sheet[sheet.max_row])
        _autosize(sheet, len(headers))

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
