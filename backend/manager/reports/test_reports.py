"""Dashboard (AD-01) and reports (AD-25, AD-26, CT-19)."""

from datetime import timedelta
from io import BytesIO

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import FarmerStatus
from manager.conftest import make_customer, make_farmer, make_market, make_order, make_product
from orders.models import Order, OrderStatus
from reviews.models import FarmerReview
from system.models import AuditAction, AuditLog

SUMMARY_URL = reverse('admin-report-summary')
EXPORT_URL = reverse('admin-report-export')


@pytest.fixture
def sales(db):
    """Two markets, two farmers; orders picked up 3 days from now."""
    customer = make_customer()
    ben_thanh, tan_dinh = make_market(), make_market('Chợ Tân Định')
    rau = make_farmer()
    trai_cay = make_farmer('f2@test.com', stall_name='Trái Cây Miền Tây')
    cabbage, mango = make_product(rau, price=15000), make_product(trai_cay, name='Xoài', price=50000)
    orders = [
        make_order(customer, rau, ben_thanh, [(cabbage, 2)], status=OrderStatus.COMPLETED),       # 30 000
        make_order(customer, rau, ben_thanh, [(cabbage, 1)], status=OrderStatus.COMPLETED),       # 15 000
        make_order(customer, trai_cay, tan_dinh, [(mango, 1)], status=OrderStatus.COMPLETED),     # 50 000
        make_order(customer, rau, ben_thanh, [(cabbage, 9)]),                                     # PLACED: no revenue
        make_order(customer, trai_cay, ben_thanh, [(mango, 1)], status=OrderStatus.DECLINED),
    ]
    FarmerReview.objects.create(order=orders[0], rating=4)
    FarmerReview.objects.create(order=orders[1], rating=5)
    pickup = orders[0].pickup_date
    return {'range': {'from': str(pickup), 'to': str(pickup)}, 'ben_thanh': ben_thanh, 'tan_dinh': tan_dinh, 'rau': rau,
            'trai_cay': trai_cay, 'pickup': pickup}


@pytest.mark.django_db
def test_dashboard(admin_client, sales):
    make_farmer('cho@test.com', stall_name='Nấm Sạch', status=FarmerStatus.PENDING)

    data = admin_client.get(reverse('admin-dashboard')).data['data']

    assert data['totals'] == {'farmers': 3, 'farmers_pending': 1, 'customers': 1, 'markets_active': 2, 'orders': 5}
    assert len(data['orders_by_day']) == 30
    assert data['orders_by_day'][-1] == {'date': timezone.localdate().isoformat(), 'count': 5}
    assert {row['status']: row['count'] for row in data['orders_by_status']} == {
        'PLACED': 1, 'ACCEPTED': 0, 'READY_FOR_PICKUP': 0, 'COMPLETED': 3,
        'CANCELLED': 0, 'DECLINED': 1, 'NO_SHOW': 0, 'EXPIRED': 0}
    assert [row['stall_name'] for row in data['pending_farmers']] == ['Nấm Sạch']


@pytest.mark.django_db
def test_orders_by_day_uses_vietnam_days(admin_client, sales):
    Order.objects.update(created_at=timezone.now() - timedelta(days=40))
    recent = Order.objects.first()
    Order.objects.filter(pk=recent.pk).update(created_at=timezone.now() - timedelta(days=1))

    data = admin_client.get(reverse('admin-dashboard')).data['data']['orders_by_day']

    assert sum(row['count'] for row in data) == 1
    assert data[-2]['count'] == 1


@pytest.mark.django_db
def test_summary_counts_only_completed_revenue(admin_client, sales):
    data = admin_client.get(SUMMARY_URL, sales['range']).data['data']

    assert {row['status']: row['count'] for row in data['orders_by_status']}['COMPLETED'] == 3
    assert data['revenue_by_market'] == [
        {'market_id': sales['tan_dinh'].pk, 'market_name': 'Chợ Tân Định', 'completed_orders': 1,
         'revenue': 50000},
        {'market_id': sales['ben_thanh'].pk, 'market_name': 'Chợ Bến Thành', 'completed_orders': 2,
         'revenue': 45000},
    ]
    assert data['top_farmers'] == [
        {'farmer_id': sales['rau'].pk, 'stall_name': 'Rau Sạch Đà Lạt', 'completed_orders': 2, 'revenue': 45000,
         'rating_avg': 4.5},
        {'farmer_id': sales['trai_cay'].pk, 'stall_name': 'Trái Cây Miền Tây', 'completed_orders': 1,
         'revenue': 50000, 'rating_avg': None},
    ]


@pytest.mark.django_db
def test_summary_filters_by_market_and_date(admin_client, sales):
    by_market = admin_client.get(SUMMARY_URL, {**sales['range'], 'market_id': sales['ben_thanh'].pk}).data['data']
    assert [row['market_name'] for row in by_market['revenue_by_market']] == ['Chợ Bến Thành']

    next_day = str(sales['pickup'] + timedelta(days=1))
    empty = admin_client.get(SUMMARY_URL, {'from': next_day, 'to': next_day}).data['data']
    assert (empty['revenue_by_market'], empty['top_farmers']) == ([], [])


@pytest.mark.django_db
@pytest.mark.parametrize('params, field', [
    ({}, 'from'),
    ({'from': '2026-01-01', 'to': 'hôm nay'}, 'to'),
    ({'from': '2026-05-01', 'to': '2026-04-01'}, 'to'),
    ({'from': '2025-01-01', 'to': '2026-01-02'}, 'to'),                   # 367 days
    ({'from': '2026-01-01', 'to': '2026-01-02', 'market_id': 999}, 'market_id'),
])
def test_summary_validation(admin_client, params, field):
    response = admin_client.get(SUMMARY_URL, params)

    assert (response.status_code, response.data['code']) == (400, 'VALIDATION_ERROR')
    assert field in response.data['errors']


@pytest.mark.django_db
def test_export_is_an_xlsx_with_three_sheets_and_is_audited(admin_client, admin_user, sales):
    # CT-19.
    response = admin_client.get(EXPORT_URL, sales['range'])

    assert response.status_code == 200
    assert response['Content-Type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml')
    day = sales['range']['from']
    assert response['Content-Disposition'] == f'attachment; filename="marketlink-report-{day}-{day}.xlsx"'
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ['Tổng quan đơn', 'Doanh thu theo chợ', 'Nông dân tích cực']
    revenue = workbook['Doanh thu theo chợ']
    assert [cell.value for cell in revenue[4]] == ['Chợ', 'Số đơn hoàn tất', 'Doanh thu (VND)']
    assert [cell.value for cell in revenue[5]] == ['Chợ Tân Định', 1, 50000]
    statuses = {row[0].value: row[1].value for row in workbook['Tổng quan đơn'].iter_rows(min_row=5)}
    assert statuses['Hoàn tất'] == 3

    entry = AuditLog.objects.get(action=AuditAction.EXPORT_DATA)
    assert (entry.user, entry.status_code, entry.details['from']) == (admin_user, 200, day)


@pytest.mark.django_db
def test_export_errors_stay_json(admin_client):
    xlsx_only = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response = admin_client.get(EXPORT_URL, HTTP_ACCEPT=xlsx_only)

    assert response.status_code == 400
    assert response.json()['code'] == 'VALIDATION_ERROR'
    assert not AuditLog.objects.filter(action=AuditAction.EXPORT_DATA).exists()


@pytest.mark.django_db
def test_customer_cannot_export(auth_client):
    assert auth_client.get(EXPORT_URL, {'from': '2026-01-01', 'to': '2026-01-02'}).status_code == 403
