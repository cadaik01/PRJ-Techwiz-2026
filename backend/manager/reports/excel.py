"""
Module: manager.reports.excel
Description: The AD-26 workbook: one sheet per AD-25 block (D-018, openpyxl).
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from orders.models import OrderStatus

VND_FORMAT = '#,##0" ₫"'
BOLD = Font(bold=True)


def _sheet(workbook, title, period, headers, rows, money_columns=()):
    sheet = workbook.create_sheet(title)
    sheet.append([title])
    sheet['A1'].font = Font(bold=True, size=13)
    sheet.append([period])
    sheet.append([])
    sheet.append(headers)
    for cell in sheet[4]:
        cell.font = BOLD
    for row in rows:
        sheet.append(row)
        for column in money_columns:
            sheet.cell(row=sheet.max_row, column=column).number_format = VND_FORMAT
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(row[index - 1])) for row in rows]) + 4
        sheet.column_dimensions[sheet.cell(row=4, column=index).column_letter].width = min(width, 45)
    return sheet


def build_workbook(report: dict, *, date_from: date, date_to: date, market_name: str | None) -> bytes:
    period = f'Ngày nhận hàng: {date_from:%d/%m/%Y} – {date_to:%d/%m/%Y}'
    if market_name:
        period += f' · Chợ: {market_name}'
    labels = dict(OrderStatus.choices)

    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, 'Tổng quan đơn', period, ['Trạng thái', 'Số đơn'],
           [[labels[row['status']], row['count']] for row in report['orders_by_status']])
    _sheet(workbook, 'Doanh thu theo chợ', period, ['Chợ', 'Số đơn hoàn tất', 'Doanh thu (VND)'],
           [[row['market_name'], row['completed_orders'], row['revenue']] for row in report['revenue_by_market']],
           money_columns=(3,))
    _sheet(workbook, 'Nông dân tích cực', period,
           ['Tên sạp', 'Số đơn hoàn tất', 'Doanh thu (VND)', 'Điểm đánh giá TB'],
           [[row['stall_name'], row['completed_orders'], row['revenue'], row['rating_avg'] or '—']
            for row in report['top_farmers']],
           money_columns=(3,))

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
